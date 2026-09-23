import pytest
import os
import io
import sys
import docx
from docx import Document
from fastapi.testclient import TestClient

backend_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

os.environ["APP_ENV"] = "test"
os.environ["JWT_SECRET"] = "test-secret-key-for-phase2-verification-2026"
os.environ["JWT_EXPIRATION_HOURS"] = "24"

from main import app
from app.services.auth_service import (
    users_db,
    _save_users_to_disk,
    hash_password,
    generate_token
)
from app.services.resume_service import (
    extract_text_from_docx,
    extract_text as extract_text_from_pdf,
    calculate_ats_score,
    generate_feedback,
    process_resume_upload
)
from app.services.skill_extractor import clean_text, parse_resume_text

client = TestClient(app)

# ==========================================
# TEST FIXTURES & BUILDERS
# ==========================================

@pytest.fixture(autouse=True)
def setup_test_users():
    test_email = "phase2-docx-tester@example.test"
    for u in list(users_db):
        if u.get("email") == test_email:
            users_db.remove(u)
            
    user = {
        "id": "test-uuid-phase2-user",
        "name": "Phase2 Tester",
        "email": test_email,
        "password": hash_password("SecurePassword123!"),
        "profile": {"name": "Phase2 Tester", "email": test_email},
        "preferences": {"targetRole": "Full Stack Engineer"},
        "scoreHistory": []
    }
    users_db.append(user)
    _save_users_to_disk()
    
    yield {"user": user}
    
    for u in list(users_db):
        if u.get("email") == test_email:
            users_db.remove(u)
    _save_users_to_disk()

def get_auth_token(user):
    token_res = generate_token(user["id"], user["email"], user["name"])
    return token_res["access_token"]

def build_minimal_pdf_bytes(lines: list) -> bytes:
    """Creates a strictly valid 1-page PDF in-memory containing the given text lines."""
    content_stream = "BT\n/F1 12 Tf\n72 720 Td\n14 TL\n"
    for line in lines:
        safe_line = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content_stream += f"({safe_line}) '\n"
    content_stream += "ET\n"
    
    stream_bytes = content_stream.encode('latin-1')
    stream_len = len(stream_bytes)
    
    body = (
        "1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        "2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        "3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>\nendobj\n"
        "4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
        f"5 0 obj\n<< /Length {stream_len} >>\nstream\n{content_stream}endstream\nendobj\n"
    )
    
    header = "%PDF-1.4\n"
    xref_pos = len(header) + len(body)
    
    # Calculate byte offsets
    o1 = len(header)
    o2 = o1 + len("1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n")
    o3 = o2 + len("2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n")
    o4 = o3 + len("3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>\nendobj\n")
    o5 = o4 + len("4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n")
    
    xref = (
        f"xref\n0 6\n"
        f"0000000000 65535 f \n"
        f"{o1:010d} 00000 n \n"
        f"{o2:010d} 00000 n \n"
        f"{o3:010d} 00000 n \n"
        f"{o4:010d} 00000 n \n"
        f"{o5:010d} 00000 n \n"
        f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    )
    
    return (header + body + xref).encode('latin-1')

def build_test_docx_bytes(
    name="Alex Mercer",
    email="alex.mercer@example.com",
    paragraphs=None,
    headings=None,
    bullets=None,
    numbers=None,
    tables=None,
    nested_tables=None
) -> bytes:
    doc = Document()
    if name:
        doc.add_paragraph(name)
    if email:
        doc.add_paragraph(email)
        
    if headings:
        for level, text in headings:
            doc.add_heading(text, level=level)
            
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
            
    if bullets:
        for b in bullets:
            doc.add_paragraph(b, style='List Bullet')
            
    if numbers:
        for n in numbers:
            doc.add_paragraph(n, style='List Number')
            
    if tables:
        for tbl_data in tables:
            rows = len(tbl_data)
            cols = len(tbl_data[0]) if rows > 0 else 0
            table = doc.add_table(rows=rows, cols=cols)
            for r_idx, row in enumerate(tbl_data):
                for c_idx, cell_value in enumerate(row):
                    table.cell(r_idx, c_idx).text = cell_value
                    
    if nested_tables:
        # Create an outer table with a nested table in one of the cells
        outer = doc.add_table(rows=2, cols=2)
        outer.cell(0, 0).text = "Outer Cell 00"
        outer.cell(0, 1).text = "Outer Cell 01"
        target_cell = outer.cell(1, 0)
        target_cell.text = "Parent Cell Header"
        inner = target_cell.add_table(rows=len(nested_tables), cols=len(nested_tables[0]))
        for r_idx, row in enumerate(nested_tables):
            for c_idx, cell_value in enumerate(row):
                inner.cell(r_idx, c_idx).text = cell_value
                
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()

# ==========================================
# 1. STRUCTURED EXTRACTION TESTS
# ==========================================

def test_1_docx_normal_paragraphs():
    docx_bytes = build_test_docx_bytes(
        paragraphs=[
            "Summary: Motivated Software Engineer with solid foundations in computer science.",
            "Passionate about building scalable web applications and clean distributed systems."
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    assert "Alex Mercer" in text
    assert "alex.mercer@example.com" in text
    assert "Motivated Software Engineer" in text
    assert "scalable web applications" in text

def test_2_docx_headings_1_to_6():
    docx_bytes = build_test_docx_bytes(
        headings=[
            (1, "Technical Skills Overview"),
            (2, "Core Frameworks"),
            (3, "Database Systems"),
            (4, "Cloud Infrastructure"),
            (5, "Methodologies"),
            (6, "Certifications")
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    assert "Technical Skills Overview" in text
    assert "Core Frameworks" in text
    assert "Database Systems" in text
    assert "Cloud Infrastructure" in text
    assert "Methodologies" in text
    assert "Certifications" in text

def test_3_docx_bullet_lists():
    docx_bytes = build_test_docx_bytes(
        bullets=[
            "Developed high-throughput APIs using FastAPI and Python.",
            "Optimized PostgreSQL database queries reducing latency by 40%.",
            "Deployed microservices on AWS with Docker and Kubernetes."
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    assert "Developed high-throughput APIs using FastAPI and Python." in text
    assert "Optimized PostgreSQL database queries" in text
    assert "Deployed microservices on AWS" in text

def test_4_docx_numbered_lists():
    docx_bytes = build_test_docx_bytes(
        numbers=[
            "B.Tech in Computer Science (CGPA: 8.8)",
            "Led full-stack web application development for client portal.",
            "Designed automated CI/CD pipeline using Jenkins and Git."
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    assert "B.Tech in Computer Science" in text
    assert "Led full-stack web application development" in text
    assert "Designed automated CI/CD pipeline" in text

def test_5_docx_standard_tables():
    docx_bytes = build_test_docx_bytes(
        tables=[
            [
                ["Category", "Proficiencies"],
                ["Languages", "Python, Java, JavaScript, TypeScript, C++"],
                ["Frameworks", "React, Node.js, FastAPI, Django, Spring Boot"],
                ["Databases", "PostgreSQL, MongoDB, MySQL"]
            ]
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    assert "Languages" in text
    assert "Python, Java, JavaScript, TypeScript, C++" in text
    assert "React, Node.js, FastAPI, Django, Spring Boot" in text
    assert "PostgreSQL, MongoDB, MySQL" in text

def test_6_docx_multiple_tables_and_nested_tables():
    docx_bytes = build_test_docx_bytes(
        tables=[
            [["Company", "Role", "Tenure"], ["TechCorp", "Software Engineer", "3+ years exp"]],
            [["Education", "Institution", "Degree"], ["State University", "Engineering College", "B.Tech"]]
        ],
        nested_tables=[
            ["Nested Skill A", "Docker"],
            ["Nested Skill B", "Kubernetes"]
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    assert "TechCorp" in text
    assert "3+ years exp" in text
    assert "State University" in text
    assert "Docker" in text
    assert "Kubernetes" in text

# ==========================================
# 2. SKILL & ENTITY EXTRACTION FIDELITY
# ==========================================

def test_7_docx_technical_skills_extraction():
    docx_bytes = build_test_docx_bytes(
        paragraphs=[
            "Technical Skills: Python, Java, React, Docker, Kubernetes, PostgreSQL, FastAPI, Git, Linux."
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    parsed = parse_resume_text(text)
    skills = [s.lower() for s in parsed["skills"]]
    assert "python" in skills
    assert "java" in skills
    assert "react" in skills
    assert "docker" in skills
    assert "kubernetes" in skills
    assert "postgresql" in skills
    assert "fastapi" in skills

def test_8_docx_abbreviations_expansion():
    docx_bytes = build_test_docx_bytes(
        paragraphs=[
            "Specialized in ML, AI, NLP, AWS cloud hosting, and DBMS design."
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    parsed = parse_resume_text(text)
    skills = [s.lower() for s in parsed["skills"]]
    keywords = [k.upper() for k in parsed["keywords"]]
    
    assert "machine learning" in skills or "ml" in [k.lower() for k in keywords]
    assert "artificial intelligence" in skills or "ai" in [k.lower() for k in keywords]
    assert "amazon web services" in skills or "aws" in [k.lower() for k in keywords]
    assert "AWS" in keywords
    assert "ML" in keywords

def test_9_docx_education_and_experience_extraction():
    docx_bytes = build_test_docx_bytes(
        paragraphs=[
            "Experience: 3+ years experience as Software Engineer at CloudSoft.",
            "Education: B.Tech in Computer Science (CGPA: 8.9)"
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    parsed = parse_resume_text(text)
    assert "3+ Years Experience" in parsed["experience"]
    assert "Software Engineering" in parsed["experience"]
    assert "B.Tech" in parsed["education"]
    assert "Computer Science" in parsed["education"]
    assert "8.9" in parsed["education"]

def test_10_docx_projects_extraction():
    docx_bytes = build_test_docx_bytes(
        paragraphs=[
            "Projects:",
            "Project: Mock Interview System with real-time feedback",
            "Project: Cloud Storage Platform using FastAPI and AWS"
        ]
    )
    text = extract_text_from_docx(docx_bytes)
    parsed = parse_resume_text(text)
    assert any("Mock Interview System" in p for p in parsed["projects"])

# ==========================================
# 3. BOUNDARY & ERROR HANDLING TESTS
# ==========================================

def test_11_empty_docx_returns_400():
    files = {'file': ('empty.docx', b'', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    res = client.post("/upload-resume", files=files)
    assert res.status_code == 400
    assert "Empty file uploaded" in res.json()["detail"]

def test_12_malformed_docx_handles_gracefully_no_500():
    # Random corrupted bytes with .docx extension
    corrupted_bytes = b"PK\x03\x04corrupted_zip_payload_which_is_not_valid_docx"
    files = {'file': ('corrupt.docx', corrupted_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    res = client.post("/upload-resume", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["data"]["name"] == "Extraction Failed"
    assert data["data"]["score"] == 0

def test_13_unsupported_encrypted_docx_no_500():
    # Simulated password-protected/encrypted OLE package header
    encrypted_bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1\x00\x00\x00\x00EncryptedPackageHeader"
    files = {'file': ('protected.docx', encrypted_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    res = client.post("/upload-resume", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["data"]["score"] == 0

# ==========================================
# 4. API CONTRACT & PERSISTENCE TESTS
# ==========================================

def test_14_docx_anonymous_upload_contract():
    docx_bytes = build_test_docx_bytes(
        name="Morgan Freeman",
        email="morgan.freeman@example.com",
        paragraphs=[
            "Skills: Python, Java, React, Docker, PostgreSQL, AWS",
            "Experience: 4 years experience as Developer",
            "Education: B.Tech in Computer Science",
            "Project: Mock Interview System"
        ]
    )
    files = {'file': ('resume_morgan.docx', docx_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    res = client.post("/upload-resume", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "filename" in data
    assert "data" in data
    ats = data["data"]
    
    # Contract validation of all required keys
    for key in ["name", "email", "skills", "experience", "education", "projects", "keywords", "score", "feedback"]:
        assert key in ats
    assert ats["score"] > 50

def test_15_docx_authenticated_upload_persists_to_user_profile(setup_test_users):
    user = setup_test_users["user"]
    token = get_auth_token(user)
    
    docx_bytes = build_test_docx_bytes(
        name="Auth Test Candidate",
        email=user["email"],
        paragraphs=[
            "Skills: Python, React, Docker, Kubernetes, SQL, AWS, Java",
            "Experience: 3 years experience as Software Engineer",
            "Education: B.E in Computer Science",
            "Project: Mock Interview System"
        ]
    )
    files = {'file': ('candidate_auth.docx', docx_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    res = client.post("/upload-resume", files=files, headers={"Authorization": f"Bearer {token}"})
    assert res.status_code == 200
    res_data = res.json()
    assert res_data["success"] is True
    
    # Verify in users_db that profile received persistent resume intelligence
    persisted_user = next(u for u in users_db if u["email"] == user["email"])
    prof = persisted_user.get("profile", {})
    assert prof.get("resumeUploaded") is True
    assert "candidate_auth.docx" in prof.get("resumeFileName", "") or "resume_" in prof.get("resumeFileName", "")
    assert prof.get("score") > 50
    assert len(prof.get("skills", [])) >= 5
    assert "Software Engineering" in prof.get("experience", "")

# ==========================================
# 5. PDF & DOCX EQUIVALENT-CONTENT PARITY TESTS
# ==========================================

def test_16_pdf_and_docx_equivalent_content_parity():
    """
    Parity Acceptance Criterion:
    Given equivalent resume content represented through supported PDF and DOCX structures:
    - normalized extracted text should be equivalent under clean_text
    - detected skills should be equivalent
    - detected keywords should be equivalent
    - education / experience / project extraction should be equivalent
    - downstream ATS score and feedback metrics should produce equivalent results under the same scoring algorithm
    """
    resume_lines = [
        "Samantha Ray",
        "samantha.ray@techdomain.org",
        "Technical Skills: Python, React, Java, Docker, PostgreSQL, AWS, ML, AI",
        "Experience: 3 years experience as Software Engineer",
        "Education: B.Tech in Computer Science (CGPA: 9.1)",
        "Projects: Project: Mock Interview System with real-time scoring"
    ]
    
    # 1. Generate equivalent PDF
    pdf_bytes = build_minimal_pdf_bytes(resume_lines)
    
    # 2. Generate equivalent DOCX
    docx_bytes = build_test_docx_bytes(
        name="Samantha Ray",
        email="samantha.ray@techdomain.org",
        paragraphs=[
            "Technical Skills: Python, React, Java, Docker, PostgreSQL, AWS, ML, AI",
            "Experience: 3 years experience as Software Engineer",
            "Education: B.Tech in Computer Science (CGPA: 9.1)",
            "Projects: Project: Mock Interview System with real-time scoring"
        ]
    )
    
    # 3. Extract text from both formats
    pdf_raw_text = extract_text_from_pdf(pdf_bytes)
    docx_raw_text = extract_text_from_docx(docx_bytes)
    
    assert "samantha ray" in clean_text(pdf_raw_text)
    assert "samantha ray" in clean_text(docx_raw_text)
    
    # 4. Parse structured ATS data
    pdf_parsed = parse_resume_text(pdf_raw_text)
    docx_parsed = parse_resume_text(docx_raw_text)
    
    # Compare skills (order-independent set equivalence)
    pdf_skills_set = set(pdf_parsed["skills"])
    docx_skills_set = set(docx_parsed["skills"])
    assert pdf_skills_set == docx_skills_set, f"Skill mismatch: PDF {pdf_skills_set} vs DOCX {docx_skills_set}"
    
    # Compare keywords (order-independent set equivalence)
    pdf_kw_set = set(pdf_parsed["keywords"])
    docx_kw_set = set(docx_parsed["keywords"])
    assert pdf_kw_set == docx_kw_set, f"Keyword mismatch: PDF {pdf_kw_set} vs DOCX {docx_kw_set}"
    
    # Compare experience, education, projects
    assert pdf_parsed["experience"] == docx_parsed["experience"]
    assert pdf_parsed["education"] == docx_parsed["education"]
    assert set(pdf_parsed["projects"]) == set(docx_parsed["projects"])
    
    # Compare ATS scoring
    pdf_score = calculate_ats_score(pdf_parsed)
    docx_score = calculate_ats_score(docx_parsed)
    assert pdf_score == docx_score
    assert pdf_score >= 80

def test_17_pdf_regression_upload():
    resume_lines = [
        "Jane Doe",
        "jane.doe@example.com",
        "Skills: Python, Java, React, SQL",
        "Experience: 2 years experience Developer",
        "Education: B.E in Information Technology"
    ]
    pdf_bytes = build_minimal_pdf_bytes(resume_lines)
    files = {'file': ('resume_regression.pdf', pdf_bytes, 'application/pdf')}
    res = client.post("/upload-resume", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "Python" in data["data"]["skills"]
    assert data["data"]["score"] > 40
